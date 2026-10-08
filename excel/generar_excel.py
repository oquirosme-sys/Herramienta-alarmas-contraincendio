# -*- coding: utf-8 -*-
"""Genera el libro Excel de la herramienta «Baterías · Alarma contra incendio» (mismas funciones que la web).

Uso (desde la carpeta del proyecto):
    python excel/generar_excel.py

Lee el catálogo de js/catalogo-base.js y escribe en excel/:
    Calculo_baterias_Alarmas_PLANTILLA.xlsx   un nivel (NIVEL 1 + FACP-01), listo para llenar
    Calculo_baterias_Alarmas_EJEMPLO.xlsx     datos del Excel original (5 paneles, 1 fuente, 2 circuitos)
Después hay que recalcular el libro (LibreOffice o abrirlo en Excel) para que queden los valores en caché.
"""
import io, json, os, re
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = os.path.join(RAIZ, 'excel')

# ------------------------------------------------------------------ catálogo
_src = io.open(os.path.join(RAIZ, 'js', 'catalogo-base.js'), encoding='utf-8').read()
_body = _src[_src.index('{', _src.index('CATALOGO_BASE')):_src.rindex('}') + 1]
CAT = json.loads(re.sub(r'^(\s*)(\w+):', r'\1"\2":', _body, flags=re.M))
FAB = {f['id']: f['nombre'] for f in CAT['fabricantes']}


def fab_key(fid):
    return FAB[fid].upper().replace('É', 'E')


def etiqueta(d):
    cod = '' if d['modelo'] in ('', '—') else d['modelo']
    t = d.get('tipo', '')
    e = cod + (' · ' if cod and t else '') + t
    return e + (' — ' if e and d['descripcion'] else '') + d['descripcion']


ETQ = {d['id']: etiqueta(d) for d in CAT['dispositivos']}
assert len(set(ETQ.values())) == len(ETQ), 'Etiquetas duplicadas: la lista desplegable no sería unívoca'


def disp(fab, tag, modelo=None):
    for d in CAT['dispositivos']:
        if d['fabricante'] == fab and d['tag'] == tag and (modelo is None or d['modelo'] == modelo):
            return d
    raise KeyError((fab, tag, modelo))


def cable_etq(modelo):
    for c in CAT['cables']:
        if c['modelo'] == modelo:
            return '%s %s (%s AWG · %s)' % (c['fabricante'], c['modelo'], c['awg'], c['listado'])
    raise KeyError(modelo)


# ------------------------------------------------------------------ estilos
ARIAL = 'Arial'
NAVY = '14263F'
F_BASE = Font(name=ARIAL, size=10)
F_B = Font(name=ARIAL, size=10, bold=True)
F_H = Font(name=ARIAL, size=10, bold=True, color='FFFFFF')
F_T = Font(name=ARIAL, size=14, bold=True, color=NAVY)
F_S = Font(name=ARIAL, size=9, italic=True, color='666666')
FILL_H = PatternFill('solid', fgColor=NAVY)
FILL_SEC = PatternFill('solid', fgColor='DCE6F2')
FILL_IN = PatternFill('solid', fgColor='E8F5E9')     # entrada manual (verde, como el Excel original)
FILL_OV = PatternFill('solid', fgColor='FFF4D6')     # reemplazo manual de corriente (amarillo)
FILL_CALC = PatternFill('solid', fgColor='F4F6F9')
FILL_RES = PatternFill('solid', fgColor='E6F4EA')
THIN = Side(style='thin', color='C9D1DB')
BORDE = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTRO = Alignment(horizontal='center', vertical='center', wrap_text=True)
WRAP = Alignment(horizontal='left', vertical='top', wrap_text=True)
DER = Alignment(horizontal='right')


def put(ws, ref, valor, font=F_BASE, fill=None, fmt=None, al=None, border=False):
    c = ws[ref]
    c.value = valor
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if al:
        c.alignment = al
    if border:
        c.border = BORDE
    return c


def inp(ws, ref, valor=None, fmt=None, fill=FILL_IN):
    return put(ws, ref, valor, F_BASE, fill, fmt, None, True)


def calc(ws, ref, f, fmt=None, bold=False):
    return put(ws, ref, f, F_B if bold else F_BASE, FILL_CALC, fmt, None, True)


def header(ws, fila, col_ini, titulos, alto=32):
    for i, t in enumerate(titulos):
        put(ws, '%s%d' % (L(col_ini + i), fila), t, F_H, FILL_H, al=CENTRO, border=True)
    ws.row_dimensions[fila].height = alto


def seccion(ws, fila, texto, c1=2, c2=2):
    for c in range(c1, c2 + 1):
        ws.cell(row=fila, column=c).fill = FILL_SEC
    put(ws, '%s%d' % (L(c1), fila), texto, F_B, FILL_SEC)


def anchos(ws, d):
    for k, v in d.items():
        ws.column_dimensions[k].width = v


def dv_lista(ws, formula, rango, estricta=True):
    dv = DataValidation(type='list', formula1=formula, allow_blank=True)
    dv.showErrorMessage = estricta
    dv.errorTitle = 'Valor no válido'
    dv.error = 'Elija un valor de la lista.'
    ws.add_data_validation(dv)
    dv.add(rango)


def pista(ws, rango, titulo, texto):
    """Ejemplo tenue: mensaje que aparece al seleccionar la celda (no ocupa la celda, que queda limpia)."""
    dv = DataValidation(allow_blank=True)
    dv.showInputMessage = True
    dv.promptTitle = titulo[:32]
    dv.prompt = texto[:255]
    ws.add_data_validation(dv)
    dv.add(rango)


def pista_texto(ws, ref, texto):
    """Línea de ejemplo en gris tenue (cursiva); no participa en ningún cálculo."""
    put(ws, ref, texto, Font(name=ARIAL, size=9, italic=True, color='9AA5B1'))


def cf_estado(ws, rango, ref):
    ws.conditional_formatting.add(rango, FormulaRule(formula=['LEFT(%s,2)="OK"' % ref], font=Font(name=ARIAL, bold=True, color='1A7F37'), fill=PatternFill('solid', bgColor='E6F4EA')))
    ws.conditional_formatting.add(rango, FormulaRule(formula=['OR(LEFT({0},5)="ERROR",LEFT({0},7)="REVISAR")'.format(ref)], font=Font(name=ARIAL, bold=True, color='C62828'), fill=PatternFill('solid', bgColor='FDE8E8')))


# ------------------------------------------------------------------ constantes de diseño
NMAX_DISP, NMAX_CAB, NMAX_BAT = 300, 60, 25
NIV = (18, 37)             # PROYECTO: lista de niveles
TAB = (42, 61)             # PROYECTO: tabla de equipos (cabecera en la fila 41)
R_D1, R_D2 = 14, 38        # hojas de equipo: filas de la tabla de dispositivos (cabecera en la 13)
R_TOT = 39
FILAS_CAIDA = (16, 45)
RESUMEN_DISP = (6, 105)

LISTA_FAB = 'OFFSET(BD_FABRICANTES!$B$6,0,0,COUNTA(BD_FABRICANTES!$B$6:$B$40),1)'
LISTA_NIVELES = 'OFFSET(PROYECTO!$B$%d,0,0,MAX(1,COUNTA(PROYECTO!$B$%d:$B$%d)),1)' % (NIV[0], NIV[0], NIV[1])
LISTA_DISP_TODOS = 'OFFSET(BD_DISPOSITIVOS!$B$6,0,0,COUNTA(BD_DISPOSITIVOS!$D$6:$D$%d),1)' % NMAX_DISP
LISTA_CABLES = 'OFFSET(BD_CABLES!$B$6,0,0,COUNTA(BD_CABLES!$D$6:$D$%d),1)' % NMAX_CAB
LISTA_EQUIPOS = 'PROYECTO!$C$%d:$C$%d' % TAB
LISTA_LAZOS = 'OFFSET(CAIDA_TENSION!$U$16,0,0,MAX(1,COUNTIF(CAIDA_TENSION!$U$16:$U$45,"?*")),1)'
TIPOS_LAZO_XL = 'SLC,NAC,IDNAC,24VDC,VOCEO,IDC'


def lista_modelos(fila):
    return ('OFFSET(BD_DISPOSITIVOS!$B$6,MATCH($B{r},BD_DISPOSITIVOS!$C$6:$C${n},0)-1,0,COUNTIF(BD_DISPOSITIVOS!$C$6:$C${n},$B{r}),1)'
            .format(r=fila, n=NMAX_DISP))


# ------------------------------------------------------------------ hojas de catálogo (ocultas)
def hoja_fabricantes(wb):
    ws = wb.create_sheet('BD_FABRICANTES')
    put(ws, 'B2', 'LISTA DE FABRICANTES / MARCAS', F_T)
    put(ws, 'B3', 'Una fila por fabricante, sin filas vacías. Debe coincidir EXACTAMENTE con la columna FABRICANTE de BD_DISPOSITIVOS (mayúsculas, sin tildes).', F_S)
    header(ws, 5, 2, ['FABRICANTE (clave)'], 20)
    for i, f in enumerate(CAT['fabricantes']):
        inp(ws, 'B%d' % (6 + i), fab_key(f['id']))
    anchos(ws, {'A': 3, 'B': 30})
    ws.sheet_state = 'hidden'


def hoja_dispositivos(wb):
    ws = wb.create_sheet('BD_DISPOSITIVOS')
    put(ws, 'B2', 'BASE DE DATOS DE DISPOSITIVOS — CORRIENTES DE ESPERA Y ALARMA (24 VDC)', F_T)
    put(ws, 'B3', 'Valores de fichas de fabricante: VERIFICAR contra la ficha vigente. Celdas verdes = editables. Mantenga los dispositivos AGRUPADOS por fabricante '
                  '(los menús filtran por bloques). Para agregar uno: inserte una fila dentro del bloque del fabricante y copie hacia abajo la fórmula de la columna B.', F_S)
    header(ws, 5, 2, ['ETIQUETA (la que aparece en las listas)', 'FABRICANTE', 'CÓDIGO / MODELO', 'TAG', 'TIPO (qué es)', 'DESCRIPCIÓN', 'I ESPERA (mA)', 'I ALARMA (mA)', 'CIRCUITO', 'FUENTE / OBSERVACIÓN'])
    for i, d in enumerate(CAT['dispositivos']):
        r = 6 + i
        vals = [fab_key(d['fabricante']), d['modelo'], d['tag'], d.get('tipo', ''), d['descripcion'], d['iEspera'], d['iAlarma'], d['circuito'], d['obs']]
        for col, v in zip('CDEFGHIJK', vals):
            inp(ws, '%s%d' % (col, r), v, '0.0##' if col in 'HI' else None)
    for r in range(6, NMAX_DISP + 1):
        cod = 'IF(OR($D{r}="",$D{r}="—"),"",$D{r})'.format(r=r)
        f = '={c}&IF(AND({c}<>"",$F{r}<>"")," · ","")&$F{r}&IF(AND({c}&$F{r}<>"",$G{r}<>"")," — ","")&$G{r}'.format(c=cod, r=r)
        calc(ws, 'B%d' % r, f)
    anchos(ws, {'A': 3, 'B': 72, 'C': 14, 'D': 18, 'E': 14, 'F': 26, 'G': 48, 'H': 12, 'I': 12, 'J': 14, 'K': 50})
    ws.freeze_panes = 'C6'
    ws.sheet_state = 'hidden'


def hoja_cables(wb):
    ws = wb.create_sheet('BD_CABLES')
    put(ws, 'B2', 'BASE DE DATOS DE CABLES FPL/FPLR/FPLP — RESISTENCIA DC NOMINAL', F_T)
    put(ws, 'B3', 'Resistencia DC nominal por CONDUCTOR a 20 °C según ficha del fabricante. El cálculo de caída usa 2 × L (ida y vuelta). Verificar ficha vigente del carrete instalado.', F_S)
    header(ws, 5, 2, ['ETIQUETA (lista)', 'FABRICANTE', 'MODELO', 'AWG', 'COND.', 'PANTALLA', 'LISTADO NEC/UL', 'R (Ω/km por conductor)', 'R (Ω/m)', 'USO TÍPICO', 'OBSERVACIÓN'])
    for i, c in enumerate(CAT['cables']):
        r = 6 + i
        for col, v in zip('CDEFGHIKL', [c['fabricante'], c['modelo'], c['awg'], c['conductores'], c['pantalla'], c['listado'], c['rKm'], c['uso'], c['obs']]):
            inp(ws, '%s%d' % (col, r), v)
    for r in range(6, NMAX_CAB + 1):
        calc(ws, 'B%d' % r, '=IF($D{r}="","",$C{r}&" "&$D{r}&" ("&$E{r}&" AWG · "&$H{r}&")")'.format(r=r))
        calc(ws, 'J%d' % r, '=IF($I{r}="","",$I{r}/1000)'.format(r=r), '0.00000')
    anchos(ws, {'A': 3, 'B': 46, 'C': 14, 'D': 14, 'E': 7, 'F': 7, 'G': 11, 'H': 18, 'I': 16, 'J': 11, 'K': 28, 'L': 50})
    ws.sheet_state = 'hidden'


def hoja_baterias(wb):
    ws = wb.create_sheet('BD_BATERIAS')
    put(ws, 'B2', 'BASE DE DATOS DE BATERÍAS 12 V SLA (2 EN SERIE = 24 V)', F_T)
    put(ws, 'B3', 'En cualquier orden: la selección toma la capacidad inmediata superior a la requerida. Verificar dimensiones vs. gabinete y cargador según ficha del panel.', F_S)
    header(ws, 5, 2, ['CAPACIDAD (Ah)', 'REF. SIMPLEX', 'REF. NOTIFIER / PS', 'REF. GENÉRICA', 'OBSERVACIÓN'], 20)
    for r in range(6, NMAX_BAT + 1):
        i = r - 6
        b = CAT['baterias'][i] if i < len(CAT['baterias']) else None
        vals = [b['ah'], b['refSimplex'], b['refNotifier'], b['refGenerica'], b['obs']] if b else [None] * 5
        for col, v in zip('BCDEF', vals):
            inp(ws, '%s%d' % (col, r), v)
    anchos(ws, {'A': 3, 'B': 16, 'C': 16, 'D': 22, 'E': 16, 'F': 46})
    ws.sheet_state = 'hidden'


# ------------------------------------------------------------------ PROYECTO
def hoja_proyecto(wb, datos, hojas_equipos):
    ws = wb.create_sheet('PROYECTO')
    put(ws, 'B2', 'PROYECTO — BATERÍAS, FUENTES AUXILIARES Y CAÍDA DE TENSIÓN (NFPA 72:2022)', F_T)
    put(ws, 'B3', 'Celdas VERDES = entrada manual · grises = cálculo automático. Cada panel/transponder y cada fuente tiene su propia hoja (ver INSTRUCCIONES).', F_S)
    seccion(ws, 4, 'INFORMACIÓN DEL PROYECTO', 2, 5)
    campos = [('NOMBRE DEL PROYECTO', datos['nombre']), ('N.º DE PROYECTO', datos['numero']), ('CLIENTE', ''), ('UBICACIÓN', ''),
              ('ELABORÓ', ''), ('REVISÓ', ''), ('EMPRESA', 'Sinergia Ingeniería'), ('FECHA', ''), ('REVISIÓN', 0),
              ('NORMATIVA', 'NFPA 72:2022 · UL 864 · NEC (NFPA 70) Art. 760')]
    for i, (k, v) in enumerate(campos):
        r = 5 + i
        put(ws, 'B%d' % r, k, F_B, FILL_CALC, border=True)
        inp(ws, 'C%d' % r, v)
        inp(ws, 'D%d' % r)
        inp(ws, 'E%d' % r)
        ws.merge_cells('C%d:E%d' % (r, r))
    ws['C12'].number_format = 'yyyy-mm-dd'
    if datos.get('pistas'):
        for fila_, tit, txt in [(5, 'Nombre del proyecto', 'Ej.: Torre Central — Fase 2'), (6, 'N.º de proyecto', 'Ej.: P-2026-014'), (7, 'Cliente', 'Ej.: Desarrolladora XYZ S.A.'),
                                (8, 'Ubicación', 'Ej.: Escazú, San José'), (9, 'Elaboró', 'Ej.: nombre de quien elabora la memoria'), (10, 'Revisó', 'Ej.: nombre de quien revisa'),
                                (12, 'Fecha', 'Ej.: 2026-10-07'), (13, 'Revisión', 'Ej.: 0, A, B…')]:
            pista(ws, 'C%d:E%d' % (fila_, fila_), tit, txt)
        pista(ws, 'B%d:B%d' % NIV, 'Niveles', 'Ej.: SÓTANO 1, NIVEL 1, NIVEL 2, AZOTEA (uno por fila)')
        pista(ws, 'B%d:B%d' % TAB, 'Hoja del equipo', 'Escriba el nombre exacto de la pestaña del panel, transponder o fuente (ej.: BAT_TRP_N2)')
    seccion(ws, 4, 'PARÁMETROS NFPA 72', 7, 9)
    pars = [('TIEMPO DE ESPERA (h)', 24, '0', '§10.6.7.2.1: 24 h'), ('TIEMPO DE ALARMA (min)', 15, '0', '5 min alarma · 15 min voceo/EVACS'),
            ('FACTOR DE SEGURIDAD', 0.2, '0%', 'Envejecimiento de baterías (fichas)'), ('TENSIÓN NOMINAL (V)', 24, '0.0', ''),
            ('FUENTE A FIN DE VIDA (%)', 0.85, '0%', 'Criterio UL 864'), ('V MÍNIMA DE DISPOSITIVO (V)', 16, '0.0', 'UL 1971/464: típico 16–33 V'),
            ('I MÁX. POR CIRCUITO NAC (A)', 2.4, '0.00', '80 % de la capacidad del NAC')]
    for i, (k, v, fmt, nota) in enumerate(pars):
        r = 5 + i
        put(ws, 'G%d' % r, k, F_B, FILL_CALC, border=True)
        inp(ws, 'H%d' % r, v, fmt)
        put(ws, 'I%d' % r, nota, F_S)
    put(ws, 'G12', 'V FUENTE A FIN DE VIDA (V)', F_B, FILL_CALC, border=True)
    calc(ws, 'H12', '=H8*H9', '0.00', True)
    put(ws, 'I12', 'Tensión nominal × % a fin de vida', F_S)

    seccion(ws, 16, 'NIVELES DEL EDIFICIO (alimentan las listas de las hojas de equipos)', 2, 3)
    header(ws, 17, 2, ['NIVEL', 'EQUIPOS EN EL NIVEL'], 32)
    for r in range(NIV[0], NIV[1] + 1):
        i = r - NIV[0]
        inp(ws, 'B%d' % r, datos['niveles'][i] if i < len(datos['niveles']) else None)
        a, b = TAB
        calc(ws, 'C%d' % r, ('=IF($B{r}="","",IF(COUNTIF($E${a}:$E${b},$B{r})=0,"(sin equipos)",INDEX($C${a}:$C${b},MATCH($B{r},$E${a}:$E${b},0))'
                             '&IF(COUNTIF($E${a}:$E${b},$B{r})>1," +"&(COUNTIF($E${a}:$E${b},$B{r})-1)&" más","")))').format(r=r, a=a, b=b))
    put(ws, 'E18', 'Para agregar un nivel: 1) escríbalo en esta lista; 2) copie una hoja de transponder (clic derecho en la pestaña →', F_S)
    put(ws, 'E19', 'Mover o copiar → Crear una copia) y renómbrela; 3) cambie TAG y NIVEL en su encabezado; 4) escriba el nombre de la', F_S)
    put(ws, 'E20', 'hoja en la columna HOJA de la tabla de abajo. (Un nivel nuevo = una hoja nueva, igual que el botón «+ Nivel» de la web.)', F_S)

    seccion(ws, 39, 'PANELES, TRANSPONDERS Y FUENTES AUXILIARES (resumen: se llena solo desde cada hoja)', 2, 11)
    put(ws, 'B40', 'Escriba en HOJA el nombre exacto de la pestaña del equipo; el resto se carga automáticamente.', F_S)
    header(ws, TAB[0] - 1, 2, ['HOJA', 'TAG', 'TIPO', 'NIVEL / UBICACIÓN', 'I ESPERA (A)', 'I ALARMA (A)', 'Ah REQUERIDO', 'BATERÍA (Ah)', 'REFERENCIA', 'ESTADO'], 32)
    celdas = [('C', '$H$3', True, None), ('D', '$H$4', True, None), ('E', '$J$3', True, None), ('F', '$Q$3', False, '0.000'), ('G', '$Q$4', False, '0.000'),
              ('H', '$Q$5', False, '0.00'), ('I', '$Q$6', False, '0.0'), ('J', '$Q$7', True, None), ('K', '$Q$8', True, None)]
    for r in range(TAB[0], TAB[1] + 1):
        i = r - TAB[0]
        inp(ws, 'B%d' % r, hojas_equipos[i] if i < len(hojas_equipos) else None)
        for col, ref, texto, fmt in celdas:
            ind = 'INDIRECT("\'"&$B{r}&"\'!{ref}")'.format(r=r, ref=ref)
            f = '=IF($B{r}="","",IFERROR({ind}{amp},"(hoja no encontrada)"))'.format(r=r, ind=ind, amp='&""' if texto else '')
            calc(ws, '%s%d' % (col, r), f, fmt)
    t = TAB[1] + 1
    put(ws, 'B%d' % t, 'TOTAL SISTEMA (paneles y transponders)', F_B, FILL_SEC)
    for c in range(3, 12):
        ws.cell(row=t, column=c).fill = FILL_SEC
    a, b = TAB
    for col in 'FGH':
        f = '=SUMIF($D${a}:$D${b},"FACP",{c}${a}:{c}${b})+SUMIF($D${a}:$D${b},"TRP",{c}${a}:{c}${b})'.format(a=a, b=b, c=col)
        put(ws, '%s%d' % (col, t), f, F_B, FILL_SEC, '0.000' if col != 'H' else '0.00')
    put(ws, 'B%d' % (t + 1), 'Cantidad de paneles / transponders:', F_BASE)
    put(ws, 'F%d' % (t + 1), '=COUNTIF($D${a}:$D${b},"FACP")+COUNTIF($D${a}:$D${b},"TRP")'.format(a=a, b=b), F_B)
    put(ws, 'B%d' % (t + 2), 'Cantidad de fuentes auxiliares:', F_BASE)
    put(ws, 'F%d' % (t + 2), '=COUNTIF($D${a}:$D${b},"FUENTE")'.format(a=a, b=b), F_B)
    put(ws, 'B%d' % (t + 3), 'Equipos por revisar:', F_BASE)
    put(ws, 'F%d' % (t + 3), '=SUMPRODUCT(--(LEFT($K${a}:$K${b},7)="REVISAR"))'.format(a=a, b=b), F_B)
    cf_estado(ws, 'K%d:K%d' % (a, b), 'K%d' % a)
    anchos(ws, {'A': 3, 'B': 34, 'C': 24, 'D': 14, 'E': 24, 'F': 14, 'G': 28, 'H': 14, 'I': 14, 'J': 34, 'K': 24})
    ws.sheet_properties.tabColor = NAVY
    return ws


# ------------------------------------------------------------------ hoja de equipo (panel / transponder / fuente)
def hoja_equipo(wb, nombre, eq):
    es_f = eq['clase'] == 'fuente'
    ws = wb.create_sheet(nombre)
    ws.sheet_properties.tabColor = '2E7D32' if es_f else '1F6FEB'
    put(ws, 'B2', '="%s — "&$H$3&" · "&$J$3&" (NFPA 72:2022 §10.6.7)"' % ('FUENTE DE PODER AUXILIAR' if es_f else 'CÁLCULO DE BATERÍAS'), F_T)
    put(ws, 'B3', 'PROYECTO:', F_B)
    put(ws, 'C3', '=PROYECTO!$C$5&""', F_BASE)
    put(ws, 'F3', 'FUENTE / TAG:' if es_f else 'PANEL / TAG:', F_B, al=DER)
    inp(ws, 'H3', eq['tag'])
    put(ws, 'I3', 'UBICACIÓN:', F_B, al=DER)
    inp(ws, 'J3', eq['nivel'])
    put(ws, 'B4', 'ELABORÓ:', F_B)
    put(ws, 'C4', '=PROYECTO!$C$9&""', F_BASE)
    put(ws, 'F4', 'TIPO:', F_B, al=DER)
    if es_f:
        calc(ws, 'H4', 'FUENTE')
    else:
        inp(ws, 'H4', eq['tipo'])
        dv_lista(ws, '"FACP,TRP"', 'H4')
    dv_lista(ws, LISTA_NIVELES, 'J3', estricta=False)

    seccion(ws, 6, 'PARÁMETROS (casilla verde vacía = se usa el valor del proyecto)', 2, 6)
    put(ws, 'D6', 'REEMPLAZO', F_B, FILL_SEC, al=CENTRO)
    put(ws, 'E6', 'VALOR USADO', F_B, FILL_SEC, al=CENTRO)
    for r, (k, ref, fmt) in zip((7, 8, 9), [('TIEMPO DE ESPERA (h)', 'H5', '0'), ('TIEMPO DE ALARMA (min)', 'H6', '0'), ('FACTOR DE SEGURIDAD', 'H7', '0%')]):
        put(ws, 'B%d' % r, k, F_B)
        inp(ws, 'D%d' % r, None, fmt)
        calc(ws, 'E%d' % r, '=IF($D${r}="",PROYECTO!${c}${n},$D${r})'.format(r=r, c=ref[0], n=ref[1:]), fmt, True)
    put(ws, 'F8', '5 min alarma · 15 min voceo/EVACS (§10.6.7.2.1)', F_S)
    if es_f:
        put(ws, 'B10', 'I MÁX. FUENTE (A) — ficha', F_B)
        inp(ws, 'D10', eq['iMax'], '0.00')
        put(ws, 'F10', 'Ej.: Notifier FCPS-24S8 = 8 A. Se limita al 80 %', F_S)
        put(ws, 'B11', 'I PROPIA DE LA FUENTE (A)', F_B)
        inp(ws, 'D11', eq['iPropia'], '0.000')
        put(ws, 'F11', 'Consumo propio del módulo según ficha; se suma en espera y en alarma', F_S)

    header(ws, 13, 2, ['FABRICANTE', 'MODELO  (código · qué es — descripción)', 'TAG', 'TIPO (qué es)', 'NIVEL / ZONA', 'CANT.', 'I ESPERA UNIT. (mA)', 'I ESPERA TOTAL (A)',
                       'I ALARMA UNIT. (mA)', 'I ALARMA TOTAL (A)', 'OBSERVACIÓN', 'LAZO / SALIDA', 'I ESPERA MANUAL (mA)', 'I ALARMA MANUAL (mA)', 'AVISO'], 44)
    BD = 'BD_DISPOSITIVOS!${c}$6:${c}$%d' % NMAX_DISP
    for r in range(R_D1, R_D2 + 1):
        i = r - R_D1
        f = eq['filas'][i] if i < len(eq['filas']) else None
        inp(ws, 'B%d' % r, fab_key(f['fab']) if f else None)
        inp(ws, 'C%d' % r, ETQ[f['disp']] if f else None)
        look = lambda col: 'INDEX({rng},MATCH($C{r},{key},0))'.format(rng=BD.format(c=col), key=BD.format(c='B'), r=r)
        calc(ws, 'D%d' % r, '=IF($C{r}="","",IFERROR({x},""))'.format(r=r, x=look('E')))
        calc(ws, 'E%d' % r, '=IF($C{r}="","",IFERROR({x},""))'.format(r=r, x=look('F')))
        inp(ws, 'F%d' % r, f['zona'] if f else None)
        inp(ws, 'G%d' % r, f['cant'] if f else None, '0')
        calc(ws, 'H%d' % r, '=IF($C{r}="","",IF($N{r}<>"",$N{r},IFERROR(IF({x}="","",{x}),"")))'.format(r=r, x=look('H')), '0.0##')
        calc(ws, 'I%d' % r, '=IF(OR($G{r}="",$H{r}=""),"",$G{r}*$H{r}/1000)'.format(r=r), '0.0000')
        calc(ws, 'J%d' % r, '=IF($C{r}="","",IF($O{r}<>"",$O{r},IFERROR(IF({x}="","",{x}),"")))'.format(r=r, x=look('I')), '0.0##')
        calc(ws, 'K%d' % r, '=IF(OR($G{r}="",$J{r}=""),"",$G{r}*$J{r}/1000)'.format(r=r), '0.0000')
        inp(ws, 'L%d' % r, f['obs'] if f else None)
        inp(ws, 'M%d' % r, f.get('circ') if f else None)
        inp(ws, 'N%d' % r, None, '0.0##', FILL_OV)
        inp(ws, 'O%d' % r, None, '0.0##', FILL_OV)
        # auxiliares (columnas ocultas): tipo de dispositivo y tipo del lazo asignado, para avisar si no corresponden
        calc(ws, 'S%d' % r, '=IF($C{r}="","",IFERROR(INDEX(BD_DISPOSITIVOS!$J$6:$J${n},MATCH($C{r},BD_DISPOSITIVOS!$B$6:$B${n},0)),""))'.format(r=r, n=NMAX_DISP))
        calc(ws, 'T%d' % r, '=IF($M{r}="","",IFERROR(INDEX(CAIDA_TENSION!$D$16:$D$45,MATCH($M{r},CAIDA_TENSION!$U$16:$U$45,0)),""))'.format(r=r))
        compat = ('IF($S{r}="SLC",$T{r}="SLC",IF($S{r}="NAC",$T{r}="NAC",IF($S{r}="IDNAC",$T{r}="IDNAC",IF($S{r}="FUENTE AUX",OR($T{r}="24VDC",$T{r}="NAC"),'
                  'IF($S{r}="VOCEO",OR($T{r}="VOCEO",$T{r}="NAC"),IF(OR($S{r}="IDC",$S{r}="ZONA"),$T{r}="IDC",TRUE))))))').format(r=r)
        calc(ws, 'P%d' % r, ('=IF(AND($C{r}="",$G{r}=""),"",IF($C{r}="","Elija el modelo",IF(ISNA(MATCH($C{r},{key},0)),"Modelo no está en el catálogo",'
                             'IF(OR($H{r}="",$J{r}=""),"Corriente sin definir en el catálogo: digítela en las columnas MANUAL",'
                             'IF(OR($G{r}="",$G{r}<=0),"Digite la cantidad",'
                             'IF(AND($M{r}<>"",$T{r}=""),"El lazo asignado ya no existe",'
                             'IF(AND($M{r}<>"",NOT({compat})),"Es de tipo "&$S{r}&" y el lazo es "&$T{r},"")))))))').format(r=r, key=BD.format(c='B'), compat=compat))
        ws['P%d' % r].font = Font(name=ARIAL, size=9, color='C62828')
    dv_lista(ws, LISTA_FAB, 'B%d:B%d' % (R_D1, R_D2))
    for r in range(R_D1, R_D2 + 1):
        dv_lista(ws, lista_modelos(r), 'C%d' % r)
    dv_lista(ws, LISTA_NIVELES, 'F%d:F%d' % (R_D1, R_D2), estricta=False)
    dv_lista(ws, LISTA_LAZOS, 'M%d:M%d' % (R_D1, R_D2))
    ws.column_dimensions['S'].hidden = True
    ws.column_dimensions['T'].hidden = True
    if eq.get('pistas'):
        pista(ws, 'G%d:G%d' % (R_D1, R_D2), 'Cantidad', 'Ej.: 30 (número de dispositivos de ese modelo en este nivel y lazo)')
        pista(ws, 'L%d:L%d' % (R_D1, R_D2), 'Observación', 'Ej.: Torre A, ala norte')
        pista(ws, 'N%d:O%d' % (R_D1, R_D2), 'Reemplazo manual', 'Solo si desea reemplazar la corriente de catálogo (mA). Si lo deja vacío se usa la del catálogo.')
        pista(ws, 'D7:D9', 'Reemplazo', 'Vacío = se usa el valor del proyecto.')
        pista(ws, 'H3', 'TAG del equipo', 'Ej.: FACP-01 o TRP-N2')
        pista_texto(ws, 'B12', 'Ejemplo de fila (no se calcula): SIMPLEX | 4098-9714 · Detector de humo — DETECTOR HUMO FOTOELÉCTRICO TrueAlarm (IDNet) | NIVEL 1 | cant. 30 | lazo FACP-01 · SLC 1.   '
                              'La primera fila suele ser el consumo propio del panel o transponder.')

    put(ws, 'G%d' % R_TOT, 'TOTALES:', F_B, al=DER)
    calc(ws, 'I%d' % R_TOT, '=SUM(I%d:I%d)' % (R_D1, R_D2), '0.0000', True)
    calc(ws, 'K%d' % R_TOT, '=SUM(K%d:K%d)' % (R_D1, R_D2), '0.0000', True)
    if es_f:
        put(ws, 'G40', 'TOTAL CON CONSUMO PROPIO:', F_B, al=DER)
        calc(ws, 'I40', '=I39+N($D$11)', '0.0000', True)
        calc(ws, 'K40', '=K39+N($D$11)', '0.0000', True)
    tot_e, tot_a = ('I40', 'K40') if es_f else ('I39', 'K39')

    seccion(ws, 41, 'CÁLCULO DE CAPACIDAD DE BATERÍA (NFPA 72:2022 §10.6.7 / UL 864)', 2, 5)
    filas = [(42, 'CORRIENTE TOTAL EN ESPERA (A)', '=' + tot_e, '0.0000'), (43, 'TIEMPO DE ESPERA (h)', '=$E$7', '0'),
             (44, 'CAPACIDAD EN ESPERA (Ah)', '=E42*E43', '0.000'), (45, 'CORRIENTE TOTAL EN ALARMA (A)', '=' + tot_a, '0.0000'),
             (46, 'TIEMPO DE ALARMA (h)', '=$E$8/60', '0.0000'), (47, 'CAPACIDAD EN ALARMA (Ah)', '=E45*E46', '0.000'),
             (48, 'CAPACIDAD CALCULADA (Ah)', '=E44+E47', '0.000'), (49, 'FACTOR DE SEGURIDAD', '=$E$9', '0%'),
             (50, 'CAPACIDAD MÍNIMA REQUERIDA (Ah)', '=E48*(1+E49)', '0.000')]
    for r, k, f, fmt in filas:
        put(ws, 'B%d' % r, k, F_B if r == 50 else F_BASE)
        calc(ws, 'E%d' % r, f, fmt, r == 50)
    put(ws, 'F44', 'Ah_espera = I_espera × t_espera', F_S)
    put(ws, 'F47', 'Ah_alarma = I_alarma × t_alarma', F_S)
    put(ws, 'F50', 'Ah_req = (Ah_espera + Ah_alarma) × (1 + FS)', F_S)
    B = 'BD_BATERIAS!$B$6:$B$%d' % NMAX_BAT
    put(ws, 'B52', 'BATERÍA SELECCIONADA (Ah, estándar):', F_B)
    put(ws, 'E52', '=IF(COUNT({b})=0,"",IF(COUNTIF({b},"<"&E50)>=COUNT({b}),"EXCEDE",SMALL({b},COUNTIF({b},"<"&E50)+1)))'.format(b=B), F_B, FILL_RES, '0.0', border=True)
    put(ws, 'B53', 'REFERENCIA SIMPLEX / NOTIFIER-PS:', F_BASE)
    calc(ws, 'E53', ('=IF(ISNUMBER(E52),IFERROR(INDEX(BD_BATERIAS!$C$6:$C${n},MATCH(E52,{b},0))&" / "&INDEX(BD_BATERIAS!$D$6:$D${n},MATCH(E52,{b},0)),"—"),"—")'
                     ).format(n=NMAX_BAT, b=B))
    put(ws, 'B54', 'ESTADO:', F_B)
    if es_f:
        estado = ('=IF(E52="EXCEDE","REVISAR: excede el catálogo de baterías",IF(LEFT($E$60,5)="ERROR","REVISAR: excede 80%% de la fuente",'
                  'IF(SUMPRODUCT(--(LEN($P$%d:$P$%d)>0),--(LEFT($P$%d:$P$%d,10)<>"Es de tipo"),--(LEFT($P$%d:$P$%d,12)<>"El lazo asig"))>0,"REVISAR: hay filas con aviso","OK")))' % ((R_D1, R_D2) * 3))
    else:
        estado = ('=IF(E52="EXCEDE","REVISAR: excede el catálogo de baterías",IF(SUMPRODUCT(--(LEN($P$%d:$P$%d)>0),--(LEFT($P$%d:$P$%d,10)<>"Es de tipo"),--(LEFT($P$%d:$P$%d,12)<>"El lazo asig"))>0,"REVISAR: hay filas con aviso","OK"))' % ((R_D1, R_D2) * 3))
    put(ws, 'E54', estado, F_B, FILL_CALC, border=True)
    put(ws, 'B55', 'Configuración: 2 baterías 12 V en serie = 24 VDC. Verificar dimensiones vs. gabinete y capacidad máxima del cargador según ficha. Si «EXCEDE»: dividir la carga o usar cargador externo listado.', F_S)
    cf_estado(ws, 'E54', 'E54')
    if es_f:
        seccion(ws, 57, 'VERIFICACIÓN DE CAPACIDAD DE LA FUENTE', 2, 5)
        put(ws, 'B58', 'I TOTAL EN ALARMA (A)')
        calc(ws, 'E58', '=K40', '0.000')
        put(ws, 'B59', 'I MÁX. PERMITIDA = 80% × I MÁX (A)')
        calc(ws, 'E59', '=IF($D$10="","",0.8*$D$10)', '0.000')
        put(ws, 'B60', 'VERIFICACIÓN DE CORRIENTE', F_B)
        calc(ws, 'E60', '=IF($D$10="","Defina I máx. de la fuente",IF(E58<=E59,"OK","ERROR: EXCEDE 80% DE LA FUENTE"))', None, True)
        cf_estado(ws, 'E60', 'E60')

    # resumen que lee la hoja PROYECTO (misma ubicación en todas las hojas de equipo)
    put(ws, 'P2', 'RESUMEN (lo lee la hoja PROYECTO — no mover)', F_H, FILL_H, al=CENTRO, border=True)
    ws.merge_cells('P2:Q2')
    for r, (k, f, fmt) in zip(range(3, 9), [('I ESPERA TOTAL (A)', '=E42', '0.0000'), ('I ALARMA TOTAL (A)', '=E45', '0.0000'), ('Ah REQUERIDO', '=E50', '0.000'),
                                            ('BATERÍA (Ah)', '=E52', '0.0'), ('REFERENCIA', '=E53', None), ('ESTADO', '=E54', None)]):
        put(ws, 'P%d' % r, k, F_B, FILL_CALC, border=True)
        calc(ws, 'Q%d' % r, f, fmt)
    anchos(ws, {'A': 3, 'B': 24, 'C': 66, 'D': 14, 'E': 24, 'F': 20, 'G': 8, 'H': 13, 'I': 13, 'J': 16, 'K': 13, 'L': 28, 'M': 26, 'N': 12, 'O': 12, 'P': 34, 'Q': 30})
    ws.freeze_panes = 'D14'
    return ws


# ------------------------------------------------------------------ caída de tensión
def suma_equipos(col, r):
    """Suma, sobre todas las hojas de equipo listadas en PROYECTO, de la columna `col` de las filas asignadas al lazo de la fila r."""
    t = []
    for k in range(TAB[0], TAB[1] + 1):
        t.append('IFERROR(IF(PROYECTO!$B${k}="",0,SUMIF(INDIRECT("\'"&PROYECTO!$B${k}&"\'!$M${a}:$M${b}"),$U{r},INDIRECT("\'"&PROYECTO!$B${k}&"\'!${c}${a}:${c}${b}"))),0)'
                 .format(k=k, a=R_D1, b=R_D2, c=col, r=r))
    return '+'.join(t)


def hoja_caida(wb, circuitos, pistas=False):
    ws = wb.create_sheet('CAIDA_TENSION')
    ws.sheet_properties.tabColor = 'B26A00'
    put(ws, 'B2', 'CAÍDA DE TENSIÓN EN LAZOS (SLC / NAC / IDNAC / 24 VDC) — MÉTODO DE CARGA CONCENTRADA', F_T)
    seccion(ws, 4, 'PARÁMETROS GLOBALES (se editan en la hoja PROYECTO)', 2, 6)
    for r, (k, f, fmt) in zip(range(5, 9), [('TENSIÓN NOMINAL (V)', '=PROYECTO!$H$8', '0.0'), ('TENSIÓN DE FUENTE A FIN DE VIDA DE BATERÍA (V)', '=PROYECTO!$H$12', '0.00'),
                                            ('TENSIÓN MÍNIMA DE DISPOSITIVO (V)', '=PROYECTO!$H$10', '0.0'), ('I MÁX. POR CIRCUITO NAC (A) — 80% de ficha', '=PROYECTO!$H$11', '0.00')]):
        put(ws, 'B%d' % r, k, F_B)
        calc(ws, 'E%d' % r, f, fmt, True)
    seccion(ws, 10, 'CÓMO SE CALCULA', 2, 12)
    put(ws, 'B11', 'Defina aquí cada lazo (SLC, NAC, IDNAC, 24 VDC, voceo, IDC) con su panel, tipo, nivel, cable y longitud. En la hoja de cada equipo, columna LAZO / SALIDA, asigne cada dispositivo a su lazo.', F_BASE)
    put(ws, 'B12', 'I lazo = Σ (cantidad × corriente de alarma de los dispositivos asignados) + OTROS (mA, solo cargas que no estén en la lista).   R lazo = 2 × L × R (Ω/m).   V disp. = V fuente − I × R lazo.', F_BASE)
    put(ws, 'B13', 'Aceptación: V disp. ≥ V mínima del dispositivo e I ≤ I máx. NAC.   El cable por defecto según la simbología CDCLH-001S: 5220UL (voceo: 5220FL).', F_S)
    header(ws, 15, 2, ['FUENTE / PANEL', 'LAZO / SALIDA', 'TIPO', 'NIVEL', 'DESCRIPCIÓN', 'DISP. (cant.)', 'OTROS (mA)', 'I DISPOSITIVOS (mA)', 'I LAZO (mA)', 'CABLE (de catálogo)', 'R (Ω/km /cond.)',
                       'LONG. IDA (m)', 'R LAZO TOTAL (Ω)', 'V FUENTE (V)', 'CAÍDA DE TENSIÓN (V)', 'V EN DISPOSITIVO (V)', '% CAÍDA', 'ESTADO', 'COMENTARIO', 'CLAVE (la usan las hojas de equipo)'], 44)
    a, b = FILAS_CAIDA
    for r in range(a, b + 1):
        c = circuitos[r - a] if r - a < len(circuitos) else None
        inp(ws, 'B%d' % r, c['fuente'] if c else None)
        inp(ws, 'C%d' % r, c['circuito'] if c else None)
        inp(ws, 'D%d' % r, c['tipo'] if c else None)
        inp(ws, 'E%d' % r, c['nivel'] if c else None)
        inp(ws, 'F%d' % r, c['desc'] if c else None)
        calc(ws, 'G%d' % r, '=IF($U{r}="","",{s})'.format(r=r, s=suma_equipos('G', r)), '0')
        inp(ws, 'H%d' % r, c['otros'] if c else None, '0.0')
        calc(ws, 'I%d' % r, '=IF($U{r}="","",1000*({s}))'.format(r=r, s=suma_equipos('K', r)), '#,##0.0')
        calc(ws, 'J%d' % r, '=IF($U{r}="","",IF(AND($G{r}=0,$H{r}=""),"",$I{r}+N($H{r})))'.format(r=r), '#,##0.0', True)
        inp(ws, 'K%d' % r, cable_etq(c['cable']) if c else None)
        calc(ws, 'L%d' % r, '=IFERROR(INDEX(BD_CABLES!$I$6:$I$%d,MATCH($K%d,BD_CABLES!$B$6:$B$%d,0)),"")' % (NMAX_CAB, r, NMAX_CAB), '0.00')
        inp(ws, 'M%d' % r, c['long'] if c else None, '0')
        calc(ws, 'N%d' % r, '=IF(OR($M{r}="",$L{r}=""),"",2*$M{r}/1000*$L{r})'.format(r=r), '0.000')
        calc(ws, 'O%d' % r, '=IF($J{r}="","",$E$6)'.format(r=r), '0.00')
        calc(ws, 'P%d' % r, '=IF(OR($J{r}="",$N{r}=""),"",$J{r}/1000*$N{r})'.format(r=r), '0.00')
        calc(ws, 'Q%d' % r, '=IF($P{r}="","",$O{r}-$P{r})'.format(r=r), '0.00', True)
        calc(ws, 'R%d' % r, '=IF($P{r}="","",$P{r}/$O{r})'.format(r=r), '0.0%')
        calc(ws, 'S%d' % r, ('=IF($J{r}="","",IF($Q{r}="","INCOMPLETO",IF(AND($Q{r}>=$E$7,$J{r}<=$E$8*1000),"OK",'
                             'IF($J{r}>$E$8*1000,"ERROR: I > I MÁX NAC","ERROR: V < V MÍN"))))').format(r=r), None, True)
        calc(ws, 'T%d' % r, ('=IF($J{r}="","",IF($S{r}="INCOMPLETO",IF($K{r}="","Seleccione el cable","Digite la longitud"),'
                             'IF($S{r}<>"OK","AUMENTE CALIBRE, REDUZCA DISTANCIA O DIVIDA EL LAZO","")))').format(r=r))
        calc(ws, 'U%d' % r, '=IF(OR($B{r}="",$C{r}=""),"",$B{r}&" · "&$C{r})'.format(r=r))
    dv_lista(ws, LISTA_EQUIPOS, 'B%d:B%d' % (a, b), estricta=False)
    dv_lista(ws, '"%s"' % TIPOS_LAZO_XL, 'D%d:D%d' % (a, b))
    dv_lista(ws, LISTA_NIVELES, 'E%d:E%d' % (a, b), estricta=False)
    dv_lista(ws, LISTA_CABLES, 'K%d:K%d' % (a, b))
    cf_estado(ws, 'S%d:S%d' % (a, b), 'S%d' % a)
    if pistas:
        pista(ws, 'C%d:C%d' % (a, b), 'Nombre del lazo', 'Ej.: SLC 1, NAC 2, 24 VDC 1')
        pista(ws, 'F%d:F%d' % (a, b), 'Descripción', 'Ej.: Detección torre A, pasillos nivel 2')
        pista(ws, 'H%d:H%d' % (a, b), 'Otros (mA)', 'Solo cargas que no estén en la lista de dispositivos del equipo (mA). Normalmente vacío.')
        pista(ws, 'M%d:M%d' % (a, b), 'Longitud de ida (m)', 'Ej.: 150 (distancia del panel al último dispositivo)')
        pista_texto(ws, 'B14', 'Ejemplo de fila (no se calcula): FACP-01 | SLC 1 | SLC | NIVEL 1 | Detección torre A | cable BELDEN 5220UL | longitud 220 m')
    put(ws, 'I%d' % (b + 1), 'TOTAL (A):', F_B, al=DER)
    calc(ws, 'J%d' % (b + 1), '=SUM(J%d:J%d)/1000' % (a, b), '0.000', True)
    put(ws, 'K%d' % (b + 1), 'Verificar contra la capacidad total del panel/fuente', F_S)
    put(ws, 'B%d' % (b + 2), 'LAZOS CON ERROR:', F_B)
    calc(ws, 'E%d' % (b + 2), '=SUMPRODUCT(--(LEFT($S$%d:$S$%d,5)="ERROR"))' % (a, b), '0', True)
    anchos(ws, {'A': 3, 'B': 24, 'C': 16, 'D': 10, 'E': 16, 'F': 28, 'G': 11, 'H': 11, 'I': 14, 'J': 13, 'K': 44, 'L': 11, 'M': 11, 'N': 11, 'O': 10, 'P': 12, 'Q': 12, 'R': 9, 'S': 24, 'T': 46, 'U': 22})
    ws.freeze_panes = 'D16'
    return ws


# ------------------------------------------------------------------ memoria de cálculo
def hoja_memoria(wb):
    ws = wb.create_sheet('MEMORIA_CALCULO')
    ws.sheet_properties.tabColor = '1A7F37'
    anchos(ws, {'A': 3, 'B': 16, 'C': 14, 'D': 24, 'E': 13, 'F': 13, 'G': 13, 'H': 13, 'I': 13, 'J': 30, 'K': 26, 'L': 28, 'M': 3})

    def texto(fila, valor, font=F_BASE, alto=None):
        ws.merge_cells('B%d:L%d' % (fila, fila))
        put(ws, 'B%d' % fila, valor, font, al=WRAP)
        ws.row_dimensions[fila].height = alto or (15 if (str(valor).startswith('=') or len(str(valor)) < 200) else 28)

    put(ws, 'B2', 'MEMORIA DE CÁLCULO — SISTEMA DE DETECCIÓN Y ALARMA CONTRA INCENDIO', F_T)
    put(ws, 'B3', 'Baterías secundarias, fuentes auxiliares y caída de tensión · NFPA 72:2022', Font(name=ARIAL, size=10, color='4A5666'))
    for i, (k, ref) in enumerate([('PROYECTO:', 'C5'), ('N.º PROYECTO:', 'C6'), ('CLIENTE:', 'C7'), ('UBICACIÓN:', 'C8'), ('ELABORÓ:', 'C9'), ('REVISÓ:', 'C10'),
                                  ('EMPRESA:', 'C11'), ('FECHA / REV.:', None), ('NORMATIVA:', 'C14')]):
        r = 5 + i
        put(ws, 'B%d' % r, k, F_B, FILL_CALC, border=True)
        ws.merge_cells('C%d:L%d' % (r, r))
        f = '=IF(PROYECTO!$%s="","",PROYECTO!$%s)' % (ref[0] + ref[1:], ref[0] + ref[1:]) if ref else '=IF(PROYECTO!$C$12="","",YEAR(PROYECTO!$C$12)&"-"&TEXT(MONTH(PROYECTO!$C$12),"00")&"-"&TEXT(DAY(PROYECTO!$C$12),"00")&" · ")&"Rev. "&PROYECTO!$C$13'
        if ref:
            f = '=PROYECTO!$%s&""' % ref
        put(ws, 'C%d' % r, f, F_BASE, border=True)
    r = 15
    texto(r, '1. CRITERIOS NORMATIVOS', F_B)
    ws['B%d' % r].fill = FILL_SEC
    lineas = [
        '• NFPA 72:2022 §10.6.7.2.1(1): la fuente secundaria debe operar el sistema 24 horas en espera y luego 5 minutos en alarma.',
        '• NFPA 72:2022 §10.6.7.2.1(2): sistemas de voceo/EVACS: 24 horas en espera y 15 minutos a carga máxima conectada.',
        '="• Parámetros de este proyecto: espera "&PROYECTO!$H$5&" h, alarma "&PROYECTO!$H$6&" min, factor de seguridad "&TEXT(PROYECTO!$H$7,"0%")&" (salvo reemplazo indicado en cada hoja de equipo)."',
        '• Factor de seguridad sobre la capacidad calculada: requisito de fichas de fabricante (Simplex, Notifier, Siemens) por envejecimiento de baterías.',
        '="• Caída de tensión: todo dispositivo debe operar dentro de su rango listado (UL 1971/UL 464, típico 16–33 V). Se usa la tensión de fuente a fin de vida de batería: "&ROUND(PROYECTO!$H$12,2)&" V ("&TEXT(PROYECTO!$H$9,"0%")&" de "&PROYECTO!$H$8&" V, criterio UL 864)."',
        '• Cableado y supervisión según NFPA 72 Cap. 12 y NEC Art. 760 (FPL/FPLR/FPLP).']
    for l in lineas:
        r += 1
        texto(r, l)
    r += 2
    texto(r, '2. METODOLOGÍA DE CÁLCULO DE BATERÍAS', F_B)
    ws['B%d' % r].fill = FILL_SEC
    for l in ['• Ah_espera = I_espera (A) × t_espera (h)   ·   Ah_alarma = I_alarma (A) × t_alarma (h)',
              '• Ah_requerido = (Ah_espera + Ah_alarma) × (1 + FS).',
              '• Corrientes unitarias tomadas del catálogo de dispositivos (fichas de fabricante); verificar contra la revisión vigente del modelo/candela/tap de planos.',
              '• Selección automática: capacidad estándar inmediata superior del catálogo de baterías (2 × 12 V en serie = 24 VDC).',
              '• Cada panel/transponder calcula su propia batería: las baterías del transponder respaldan solo su carga local.']:
        r += 1
        texto(r, l)
    r += 2
    texto(r, '3. METODOLOGÍA DE CAÍDA DE TENSIÓN (CARGA CONCENTRADA)', F_B)
    ws['B%d' % r].fill = FILL_SEC
    for l in ['• R_lazo = 2 × L × r, con L = longitud de ida (m) y r = resistencia por conductor (Ω/m) de la ficha del fabricante.',
              '="• V_dispositivo = V_fuente − I_circuito × R_lazo, con V_fuente = "&ROUND(PROYECTO!$H$12,2)&" V."',
              '="• Aceptación: V_dispositivo ≥ "&ROUND(PROYECTO!$H$10,1)&" V e I_circuito ≤ "&ROUND(PROYECTO!$H$11,2)&" A (80 % de la corriente máxima del NAC/fuente)."',
              '• El método de carga concentrada (toda la carga al final del circuito) es el más conservador aceptado por NFPA 72 (Anexo A).']:
        r += 1
        texto(r, l)
    r += 2
    texto(r, '4. RESULTADOS', F_B)
    ws['B%d' % r].fill = FILL_SEC
    t = TAB[1] + 1
    kp = [('Ah requerido total del sistema (paneles):', '=PROYECTO!$H$%d' % t, '0.00'), ('Cantidad de paneles / transponders:', '=PROYECTO!$F$%d' % (t + 1), '0'),
          ('Cantidad de fuentes auxiliares:', '=PROYECTO!$F$%d' % (t + 2), '0'), ('Equipos por revisar:', '=PROYECTO!$F$%d' % (t + 3), '0'),
          ('Lazos con error de caída de tensión:', '=CAIDA_TENSION!$E$%d' % (FILAS_CAIDA[1] + 2), '0')]
    for k, f, fmt in kp:
        r += 1
        ws.merge_cells('B%d:E%d' % (r, r))
        put(ws, 'B%d' % r, k, F_B)
        put(ws, 'F%d' % r, f, F_B, FILL_RES, fmt, border=True)
    r += 2
    put(ws, 'B%d' % r, '4.1 Resumen de baterías por panel, transponder y fuente auxiliar', F_B)
    r += 1
    header(ws, r, 2, ['TAG', 'TIPO', 'UBICACIÓN', 'I ESPERA (A)', 'I ALARMA (A)', 'Ah REQUERIDO', 'BATERÍA (Ah)', 'REFERENCIA', '', 'ESTADO'], 30)
    ws.merge_cells('I%d:J%d' % (r, r))
    a, b = TAB
    primero = r + 1
    for k in range(a, b + 1):
        r += 1
        m = [('B', 'C', 0, None), ('C', 'D', 0, None), ('D', 'E', 0, None), ('E', 'F', 1, '0.000'), ('F', 'G', 1, '0.000'), ('G', 'H', 1, '0.00'), ('H', 'I', 1, '0.0'),
             ('I', 'J', 0, None), ('J', None, 0, None), ('K', 'K', 0, None)]
        for dest, orig, num, fmt in m:
            f = '=IF(PROYECTO!$B{k}="","",PROYECTO!${o}{k})'.format(k=k, o=orig) if orig else None
            put(ws, '%s%d' % (dest, r), f, F_BASE, None, fmt, border=True)
        ws.merge_cells('I%d:J%d' % (r, r))
    cf_estado(ws, 'K%d:K%d' % (primero, r), 'K%d' % primero)
    r += 2
    put(ws, 'B%d' % r, '4.2 Caída de tensión por lazo', F_B)
    r += 1
    header(ws, r, 2, ['FUENTE', 'LAZO', 'NIVEL / DESCRIPCIÓN', 'DISP.', 'I (mA)', 'L (m)', 'R LAZO (Ω)', 'CAÍDA (V)', 'V DISP. (V)', 'ESTADO', 'COMENTARIO'], 30)
    a, b = FILAS_CAIDA
    primero = r + 1
    for k in range(a, b + 1):
        r += 1
        m = [('B', 'CAIDA_TENSION!$B{k}', None), ('C', 'CAIDA_TENSION!$C{k}', None),
             ('D', 'CAIDA_TENSION!$E{k}&IF(AND(CAIDA_TENSION!$E{k}<>"",CAIDA_TENSION!$F{k}<>"")," — ","")&CAIDA_TENSION!$F{k}', None),
             ('E', 'CAIDA_TENSION!$G{k}', '0'), ('F', 'CAIDA_TENSION!$J{k}', '#,##0'), ('G', 'CAIDA_TENSION!$M{k}', '0'), ('H', 'CAIDA_TENSION!$N{k}', '0.00'),
             ('I', 'CAIDA_TENSION!$P{k}', '0.00'), ('J', 'CAIDA_TENSION!$Q{k}', '0.00'), ('K', 'CAIDA_TENSION!$S{k}', None), ('L', 'CAIDA_TENSION!$T{k}', None)]
        for dest, orig, fmt in m:
            f = '=IF(CAIDA_TENSION!$U{k}="","",{o})'.format(k=k, o=orig.format(k=k))
            put(ws, '%s%d' % (dest, r), f, F_BASE, None, fmt, border=True)
    cf_estado(ws, 'K%d:K%d' % (primero, r), 'K%d' % primero)
    r += 2
    texto(r, 'NOTA: Memoria de cálculo de ingeniería. Los valores de corriente de dispositivos y resistencia de cables provienen de fichas técnicas y deben verificarse contra la revisión vigente '
             'antes de construcción. Aprobación final: AHJ (Ingeniería de Bomberos de Costa Rica).', F_S, 30)
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = 'B2:L%d' % r
    return ws


# ------------------------------------------------------------------ resumen de dispositivos
def hoja_resumen_disp(wb, hojas):
    ws = wb.create_sheet('RESUMEN_DISPOSITIVOS')
    put(ws, 'B2', 'RESUMEN DE DISPOSITIVOS POR EQUIPO', F_T)
    put(ws, 'B3', 'Cuenta cuántos dispositivos de cada modelo hay en cada hoja de equipo. Escriba en la fila 5 el nombre de cada hoja (hasta 12 equipos). Use el filtro de TOTAL para ocultar los modelos sin uso.', F_S)
    put(ws, 'B4', 'TAG →', F_B, al=DER)
    header(ws, 5, 2, ['DISPOSITIVO (código · tipo — descripción)'] + [''] * 12 + ['TOTAL'], 30)
    for k in range(12):
        col = L(3 + k)
        inp(ws, '%s5' % col, hojas[k] if k < len(hojas) else None).alignment = CENTRO
        calc(ws, '%s4' % col, '=IF({c}5="","",IFERROR(INDIRECT("\'"&{c}5&"\'!$H$3")&"",""))'.format(c=col))
        ws['%s4' % col].font = F_B
        ws['%s4' % col].alignment = CENTRO
    a, b = RESUMEN_DISP
    for r in range(a, b + 1):
        calc(ws, 'B%d' % r, '=IF(BD_DISPOSITIVOS!$B{r}="","",BD_DISPOSITIVOS!$B{r})'.format(r=r))
        for k in range(12):
            col = L(3 + k)
            s = ('SUMIF(INDIRECT("\'"&{c}$5&"\'!$C${d1}:$C${d2}"),$B{r},INDIRECT("\'"&{c}$5&"\'!$G${d1}:$G${d2}"))').format(c=col, r=r, d1=R_D1, d2=R_D2)
            calc(ws, '%s%d' % (col, r), '=IF(OR($B{r}="",{c}$5=""),"",IF(IFERROR({s},0)=0,"",{s}))'.format(r=r, c=col, s=s), '0')
        calc(ws, 'O%d' % r, '=IF($B{r}="","",IF(SUM($C{r}:$N{r})=0,"",SUM($C{r}:$N{r})))'.format(r=r), '0', True)
    ws.auto_filter.ref = 'B5:O%d' % b
    anchos(ws, {'A': 3, 'B': 76, 'O': 10, **{L(3 + k): 12 for k in range(12)}})
    ws.freeze_panes = 'C6'
    return ws


# ------------------------------------------------------------------ instrucciones
def hoja_instrucciones(wb):
    ws = wb.create_sheet('INSTRUCCIONES')
    put(ws, 'B2', 'INSTRUCCIONES DE USO', F_T)
    pasos = [
        ('LEYENDA', None),
        ('Verde', 'Entrada manual (escriba o elija de la lista). En la plantilla están limpias; al seleccionar una casilla aparece un ejemplo tenue.'), ('Amarillo', 'Reemplazo manual de una corriente unitaria de catálogo (columnas «MANUAL» de las hojas de equipo).'),
        ('Gris', 'Cálculo automático: no escribir encima.'), ('', ''),
        ('FLUJO DE TRABAJO', None),
        ('1', 'En PROYECTO: complete los datos, revise los parámetros NFPA 72 y escriba los NIVELES del edificio (se empieza con uno; los demás se agregan como se indica abajo).'),
        ('2', 'En cada hoja de equipo (BAT_FACP_…, BAT_TRP_…, FUENTE_…): elija FABRICANTE → MODELO (la lista muestra «código · qué es — descripción») y digite NIVEL/ZONA, LAZO / SALIDA (la lista sale de la hoja CAIDA_TENSION) y CANTIDAD. La primera fila es el consumo propio del panel/transponder.'),
        ('3', 'Cada equipo calcula su propia batería y selecciona la capacidad estándar inmediata superior del catálogo (2 × 12 V en serie = 24 VDC). Si excede la mayor batería, el estado dice REVISAR.'),
        ('4', 'En CAIDA_TENSION defina primero los lazos (panel, nombre, tipo SLC/NAC/IDNAC/24 VDC/voceo/IDC, nivel, cable y longitud de ida); la corriente de cada lazo sale de los dispositivos asignados en las hojas de equipo (más «OTROS» en mA).'),
        ('5', 'MEMORIA_CALCULO reúne todos los resultados y está lista para imprimir o guardar como PDF (Archivo → Imprimir). RESUMEN_DISPOSITIVOS cuenta los modelos por equipo.'),
        ('', ''),
        ('AGREGAR UN NIVEL (equivale al botón «+ Nivel» de la web)', None),
        ('a', 'En PROYECTO escriba el nuevo nivel en la lista NIVELES.'),
        ('b', 'Clic derecho en la pestaña de un transponder → Mover o copiar → «Crear una copia» → renómbrela (p. ej. BAT_TRP_N2).'),
        ('c', 'En la copia cambie TAG (H3) y UBICACIÓN (J3) y borre las filas de dispositivos que no apliquen.'),
        ('d', 'En la tabla de PROYECTO (columna HOJA) escriba el nombre exacto de la pestaña nueva; el resumen se completa solo. Para verla en RESUMEN_DISPOSITIVOS escríbala también en la fila 5 de esa hoja.'),
        ('', ''),
        ('ADMINISTRADOR (catálogos)', None),
        ('•', 'Las hojas BD_FABRICANTES, BD_DISPOSITIVOS, BD_CABLES y BD_BATERIAS están OCULTAS. Para editarlas: clic derecho en cualquier pestaña → Mostrar…'),
        ('•', 'Agregar marca/fabricante: escríbalo en BD_FABRICANTES (MAYÚSCULAS, sin tildes) y use esa misma clave en la columna FABRICANTE de sus dispositivos.'),
        ('•', 'Agregar dispositivo: inserte una fila DENTRO del bloque de su fabricante en BD_DISPOSITIVOS y copie hacia abajo la fórmula de la columna B (ETIQUETA). Complete código, TAG, tipo, descripción y corrientes.'),
        ('•', 'Agregar cable: nueva fila en BD_CABLES (la etiqueta y R en Ω/m se calculan solos). Agregar batería: nueva fila en BD_BATERIAS (cualquier orden).'),
        ('•', 'Para que solo el administrador pueda mostrar esas hojas: Revisar → Proteger libro (estructura) con contraseña. Mientras esté protegido tampoco se pueden copiar hojas.'),
        ('', ''),
        ('IMPORTANTE', None),
        ('•', 'Todas las corrientes y resistencias provienen de fichas de fabricante: verificar contra la revisión vigente del modelo/candela/tap de planos. Aprobación final: AHJ.'),
        ('•', 'El catálogo fue validado contra la simbología CDCLH-001S (Circuito S.A., versiones Simplex y Notifier). Los dispositivos agregados sin corriente conocida aparecen con AVISO «Corriente sin definir»: digite la corriente en las columnas MANUAL (amarillas) o complete el catálogo; mientras tanto el equipo queda en REVISAR.'),
        ('•', 'Este libro es una copia local de la herramienta web; ambos usan el mismo catálogo base, los mismos lazos por dispositivo y las mismas fórmulas, pero los datos no se sincronizan entre sí.'),
    ]
    r = 3
    for k, v in pasos:
        r += 1
        if v is None:
            seccion(ws, r, k, 2, 3)
            continue
        c = put(ws, 'B%d' % r, k, F_B, al=Alignment(horizontal='center', vertical='top'))
        put(ws, 'C%d' % r, v, F_BASE, al=WRAP)
        ws.row_dimensions[r].height = 15 if len(v) < 120 else 30
        if k == 'Verde':
            c.fill = FILL_IN
        elif k == 'Amarillo':
            c.fill = FILL_OV
        elif k == 'Gris':
            c.fill = FILL_CALC
    anchos(ws, {'A': 3, 'B': 12, 'C': 140})
    return ws


# ------------------------------------------------------------------ libro
def construir(ruta, ejemplo):
    wb = Workbook()
    wb.remove(wb.active)
    sx = 'simplex'

    def fila(fab, tag, modelo, zona, cant, obs=''):
        return {'fab': fab, 'disp': disp(fab, tag, modelo)['id'], 'zona': zona, 'cant': cant, 'obs': obs}

    if ejemplo:
        datos = {'nombre': 'Proyecto de ejemplo (Excel original)', 'numero': '', 'niveles': ['SÓTANO 2', 'NIVEL 1', 'NIVEL 3', 'NIVEL 6', 'NIVEL 9']}
        equipos = [
            ('BAT_FACP_01', {'clase': 'panel', 'tag': 'FACP-01', 'tipo': 'FACP', 'nivel': 'NIVEL 1', 'filas': [fila(sx, 'FACP', '4010ES', 'NIVEL 1', 14, '(consumo propio del equipo)')]}),
            ('BAT_TRP_S2', {'clase': 'panel', 'tag': 'TRP-S2', 'tipo': 'TRP', 'nivel': 'SÓTANO 2', 'filas': [fila(sx, 'TRP', '4100-9600', 'SÓTANO 2', 1, '(consumo propio del equipo)')]}),
            ('BAT_TRP_N3', {'clase': 'panel', 'tag': 'TRP-N3', 'tipo': 'TRP', 'nivel': 'NIVEL 3', 'filas': [fila(sx, 'TRP', '4100-9600', 'NIVEL 3', 1, '(consumo propio del equipo)')]}),
            ('BAT_TRP_N6', {'clase': 'panel', 'tag': 'TRP-N6', 'tipo': 'TRP', 'nivel': 'NIVEL 6', 'filas': [fila(sx, 'TRP', '4100-9600', 'NIVEL 6', 1, '(consumo propio del equipo)')]}),
            ('BAT_TRP_N9', {'clase': 'panel', 'tag': 'TRP-N9', 'tipo': 'TRP', 'nivel': 'NIVEL 9', 'filas': [fila(sx, 'TRP', '4100-9600', 'NIVEL 9', 1, '(consumo propio del equipo)')]}),
            ('FUENTE_RPS_01', {'clase': 'fuente', 'tag': 'RPS-01', 'nivel': 'NIVEL 1', 'iMax': 8, 'iPropia': 0.145,
                               'filas': [fila(sx, 'ST-P 110cd', None, 'HABITACIONES N2', 20, '(fila de ejemplo — reemplazar)')]}),
        ]
        # mismas corrientes del Excel original: 5 × 74 mA + 8 × 84 mA = 1042 mA y 1500 mA (se capturan en «otros»)
        circ = [{'fuente': 'FACP-01', 'circuito': 'NAC 1', 'tipo': 'NAC', 'nivel': 'NIVEL 1', 'desc': 'Parqueo — ejemplo', 'otros': 1042, 'cable': '5220UL', 'long': 150},
                {'fuente': 'FACP-01', 'circuito': 'SLC 1', 'tipo': 'SLC', 'nivel': 'NIVEL 1', 'desc': 'Lazo detección N1-N3 — ejemplo', 'otros': 1500, 'cable': '5220UL', 'long': 350}]
    else:
        # plantilla limpia: sin datos; los ejemplos aparecen como mensajes tenues al seleccionar cada casilla
        datos = {'nombre': '', 'numero': '', 'niveles': ['NIVEL 1'], 'pistas': True}
        equipos = [('BAT_FACP_01', {'clase': 'panel', 'tag': 'FACP-01', 'tipo': 'FACP', 'nivel': 'NIVEL 1', 'filas': [], 'pistas': True})]
        circ = []

    hoja_proyecto(wb, datos, [n for n, _ in equipos])
    for nombre, eq in equipos:
        hoja_equipo(wb, nombre, eq)
    hoja_caida(wb, circ, pistas=not ejemplo)
    hoja_memoria(wb)
    hoja_resumen_disp(wb, [n for n, _ in equipos])
    hoja_instrucciones(wb)
    hoja_fabricantes(wb)
    hoja_dispositivos(wb)
    hoja_cables(wb)
    hoja_baterias(wb)
    for ws in wb.worksheets:
        if ws.sheet_state == 'visible' and ws.title != 'MEMORIA_CALCULO':
            ws.page_setup.orientation = 'landscape'
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    wb.calculation.fullCalcOnLoad = True
    wb.active = 0
    wb.save(ruta)
    return ruta


if __name__ == '__main__':
    os.makedirs(SALIDA, exist_ok=True)
    for nombre, ej in [('Calculo_baterias_Alarmas_PLANTILLA.xlsx', False), ('Calculo_baterias_Alarmas_EJEMPLO.xlsx', True)]:
        print(construir(os.path.join(SALIDA, nombre), ej))
